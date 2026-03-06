import os
import sys
import time
import logging
import argparse
from datetime import datetime
from tqdm import tqdm
import yaml

# 第三方库导入
from langchain_community.document_loaders import PyMuPDFLoader, Docx2txtLoader, TextLoader
from langchain_openai import ChatOpenAI

# ================= 全局初始化 =================
def parse_args():
    """解析命令行参数"""
    parser = argparse.ArgumentParser(description='硬件数据手册解析归档 Agent')
    parser.add_argument(
        '-f', '--file', 
        type=str, 
        help='指定要解析的文档文件（支持相对input_dir的路径/绝对路径），优先级高于配置文件'
    )
    parser.add_argument(
        '-o', '--output',
        type=str,
        help='指定输出文档的基础名称（不含后缀），优先级高于配置文件'
    )
    return parser.parse_args()

def setup_logging(config):
    """配置日志系统"""
    log_dir = config['logging']['log_dir']
    if not os.path.exists(log_dir):
        os.makedirs(log_dir)
    
    log_filename = os.path.join(log_dir, f"agent_{datetime.now().strftime('%Y%m%d')}.log")
    
    logging.basicConfig(
        level=getattr(logging, config['logging']['level']),
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(log_filename, encoding='utf-8'),
            logging.StreamHandler(sys.stdout)
        ]
    )
    return logging.getLogger(__name__)

def load_config(config_path="config.yaml"):
    """加载配置文件"""
    if not os.path.exists(config_path):
        print(f"❌ 错误：找不到配置文件 {config_path}！")
        print(f"   请复制 config.yaml.example 为 config.yaml 并填入配置。")
        sys.exit(1)
    
    with open(config_path, 'r', encoding='utf-8') as f:
        try:
            return yaml.safe_load(f)
        except yaml.YAMLError as e:
            print(f"❌ 配置文件格式错误: {e}")
            sys.exit(1)

# ================= 智能文档加载器 =================
def load_document_smart(file_path, logger):
    """根据文件后缀智能选择加载器"""
    if not os.path.exists(file_path):
        logger.error(f"找不到输入文件: {file_path}")
        return None

    ext = os.path.splitext(file_path)[1].lower()
    
    loader_map = {
        '.pdf': PyMuPDFLoader,
        '.docx': Docx2txtLoader,
        '.txt': TextLoader
    }

    if ext not in loader_map:
        logger.error(f"不支持的文件格式: {ext}，当前支持: {list(loader_map.keys())}")
        return None

    try:
        logger.info(f"正在加载 {ext.upper()} 文档: {os.path.basename(file_path)}")
        loader_class = loader_map[ext]
        
        if ext == '.txt':
            loader = loader_class(file_path, encoding='utf-8')
        else:
            loader = loader_class(file_path)
            
        documents = loader.load()
        full_text = "\n".join([doc.page_content for doc in documents])
        logger.info(f"文档加载成功，共 {len(documents)} 页/段")
        return full_text
    except Exception as e:
        logger.error(f"加载文档失败: {str(e)}", exc_info=True)
        return None

# ================= 表格解析辅助函数 =================
def extract_markdown_table(md_content: str, chapter_title: str, logger):
    """
    从Markdown内容中提取指定章节的表格，返回pandas.DataFrame
    适配AI生成的固定章节结构，异常场景自动兜底，新手友好
    :param md_content: AI生成的完整Markdown归档内容
    :param chapter_title: 要提取的章节标题（如"链路层数据帧结构"）
    :param logger: 日志对象
    :return: 提取成功返回DataFrame，失败返回None
    """
    try:
        # 1. 按换行拆分内容，逐行匹配目标章节
        lines = md_content.split('\n')
        chapter_found = False
        table_lines = []

        # 2. 匹配目标章节，捕获章节后的Markdown表格
        for line in lines:
            # 匹配目标章节标题（兼容# 层级、前后空格）
            if not chapter_found and chapter_title in line and line.startswith('#'):
                chapter_found = True
                continue
            # 找到章节后，开始捕获表格（以|开头的行是Markdown表格行）
            if chapter_found:
                # 遇到下一个章节标题，停止捕获（表格结束）
                if line.startswith('#'):
                    break
                # 只保留表格行
                if line.strip().startswith('|') and line.strip().endswith('|'):
                    table_lines.append(line.strip())
        
        # 兜底：未捕获到表格行
        if len(table_lines) < 2:
            logger.warning(f"⚠️ 章节【{chapter_title}】未找到有效表格，跳过")
            return None

        # 3. 清洗Markdown表格，转为标准CSV格式
        cleaned_rows = []
        for row in table_lines:
            # 去掉首尾的|，按|拆分单元格，去除每个单元格前后空格
            cells = [cell.strip() for cell in row[1:-1].split('|')]
            # 过滤Markdown表格的分隔行（---|---格式）
            if not all(cell.replace('-', '').strip() == '' for cell in cells):
                cleaned_rows.append(cells)
        
        # 兜底：清洗后无有效数据行
        if len(cleaned_rows) < 1:
            logger.warning(f"⚠️ 章节【{chapter_title}】表格清洗后无有效数据，跳过")
            return None

        # 4. 转为DataFrame返回
        import pandas as pd
        df = pd.DataFrame(cleaned_rows[1:], columns=cleaned_rows[0])
        logger.info(f"✅ 成功提取【{chapter_title}】表格，共{len(df)}条数据")
        return df

    except Exception as e:
        logger.error(f"❌ 提取【{chapter_title}】表格失败: {str(e)}", exc_info=True)
        return None
# ================= 多格式归档保存器 =================
def save_archive_multi_format(content, base_name, config, logger):
    """保存为 Markdown/Word/Excel 多种格式"""
    archive_root = config['archive'].get('archive_root_dir', 'output_archive')
    if config['archive'].get('use_timestamp_folder', True):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        target_dir = os.path.join(archive_root, f"{base_name}_{timestamp}")
    else:
        target_dir = archive_root

    if not os.path.exists(target_dir):
        os.makedirs(target_dir)
    
    saved_files = []
    formats = config['document'].get('output_formats', ['markdown'])

    # 保存 Markdown
    if "markdown" in formats:
        md_path = os.path.join(target_dir, f"{base_name}.md")
        try:
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(content)
            logger.info(f"✅ Markdown 已保存: {md_path}")
            saved_files.append(md_path)
        except Exception as e:
            logger.error(f"保存 Markdown 失败: {e}")

    # 保存 Word (.docx)
    if "word" in formats:
        docx_path = os.path.join(target_dir, f"{base_name}.docx")
        try:
            from docx import Document
            doc = Document()
            
            title = doc.add_heading(base_name, 0)
            title.alignment = 1
            
            for line in content.split('\n'):
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.strip() == '':
                    continue
                else:
                    p = doc.add_paragraph(line)
                    p.paragraph_format.line_spacing = 1.5
            
            doc.save(docx_path)
            logger.info(f"✅ Word 文档已保存: {docx_path}")
            saved_files.append(docx_path)
        except ImportError:
            logger.warning("⚠️ 未安装 python-docx 库，跳过 Word 生成。请运行: pip install python-docx")
        except Exception as e:
            logger.error(f"保存 Word 失败: {e}")

    # 保存 Excel
       # 保存 Excel
    if "excel" in formats:
        xlsx_path = os.path.join(target_dir, f"{base_name}_指令集_错误码.xlsx")
        try:
            import pandas as pd
            
            with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
                # 1. 保留说明Sheet，补充完整生成信息
                df_info = pd.DataFrame({'说明': [
                    '本文件由硬件文档Agent自动生成', 
                    f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
                    f'源文档: {base_name}',
                    '包含内容：数据帧结构、核心指令集、错误码汇总'
                ]})
                df_info.to_excel(writer, sheet_name='文档说明', index=False)

                # 2. 批量提取核心表格，写入对应Sheet
                # 定义要提取的章节与Sheet名映射（与AI生成的固定结构完全匹配）
                table_map = {
                    "链路层数据帧结构": "数据帧结构",
                    "核心指令集汇总": "核心指令集",
                    "错误码汇总表": "错误码汇总"
                }

                for chapter_title, sheet_name in table_map.items():
                    table_df = extract_markdown_table(content, chapter_title, logger)
                    if table_df is not None and not table_df.empty:
                        table_df.to_excel(writer, sheet_name=sheet_name, index=False)

            logger.info(f"✅ Excel 已保存: {xlsx_path}")
            saved_files.append(xlsx_path)
        except ImportError:
            logger.warning("⚠️ 未安装 pandas/openpyxl 库，跳过 Excel 生成。请运行: pip install pandas openpyxl")
        except Exception as e:
            logger.error(f"保存 Excel 失败: {e}")

        # 保存 CSV
    if "csv" in formats:
        try:
            import pandas as pd
            # 定义要提取的章节与CSV文件名映射
            table_map = {
                "链路层数据帧结构": "数据帧结构",
                "核心指令集汇总": "核心指令集",
                "错误码汇总表": "错误码汇总"
            }

            csv_saved_count = 0
            for chapter_title, file_suffix in table_map.items():
                csv_path = os.path.join(target_dir, f"{base_name}_{file_suffix}.csv")
                table_df = extract_markdown_table(content, chapter_title, logger)
                if table_df is not None and not table_df.empty:
                    # 用utf-8-sig编码，兼容Windows Excel打开不乱码
                    table_df.to_csv(csv_path, index=False, encoding='utf-8-sig')
                    logger.info(f"✅ CSV 已保存: {csv_path}")
                    saved_files.append(csv_path)
                    csv_saved_count += 1
            
            if csv_saved_count == 0:
                logger.warning("⚠️ 未提取到有效表格，无CSV文件生成")

        except ImportError:
            logger.warning("⚠️ 未安装 pandas 库，跳过 CSV 生成。请运行: pip install pandas")
        except Exception as e:
            logger.error(f"保存 CSV 失败: {e}")        

    return saved_files

# ================= 核心业务逻辑 =================
def process_document(config, input_file_path, output_base_name, logger):
    """主处理流程"""
    start_time = time.time()

    try:
        with tqdm(total=100, desc="整体进度", bar_format='{l_bar}{bar}| {n_fmt}%') as pbar:
            
            # 1. 智能加载文档
            pbar.set_description("正在读取文档")
            full_doc_text = load_document_smart(input_file_path, logger)
            if full_doc_text is None:
                return False
            pbar.update(30)

            # 2. 初始化 LLM
            pbar.set_description("正在初始化 AI 模型")
            logger.info("正在连接 AI 服务...")
            try:
                llm = ChatOpenAI(
                    model=config['llm']['model_name'],
                    api_key=config['llm']['api_key'],
                    base_url=config['llm']['base_url'],
                    temperature=config['llm']['temperature']
                )
            except Exception as e:
                logger.error(f"AI 服务初始化失败: {e}")
                return False
            pbar.update(10)

            # 3. 构造动态Prompt（标题和输入文档匹配）
            pbar.set_description("正在构造分析请求")
            archive_prompt = f"""
            你是一名资深的嵌入式通信协议工程师。请基于提供的文档生成一份严谨的归档文档。
            【强制结构】
            # {output_base_name}
            ## 1. 文档概述
            ## 2. 物理层通信参数
            ## 3. 链路层数据帧结构 (表格)
            ## 4. 核心指令集汇总 (表格)
            ## 5. 校验算法定义
            ## 6. 模块全生命周期工作时序
            ## 7. 特殊时序与约束标记 (用【⚠️】开头)
            ## 8. 典型交互示例 (3个)
            ## 9. 错误码汇总表 (表格)
            【参考文档】
            {full_doc_text[:30000]}
            """
            pbar.update(10)

            # 4. 调用 AI 生成
            pbar.set_description("AI 正在分析文档")
            logger.info("正在调用 AI 生成归档内容...")
            ai_start_time = time.time()
            
            ai_response = llm.invoke(archive_prompt)
            
            ai_duration = time.time() - ai_start_time
            logger.info(f"AI 生成完成，耗时: {ai_duration:.2f}秒")
            pbar.update(40)

            # 5. 多格式保存
            pbar.set_description("正在保存归档文件")
            saved_files = save_archive_multi_format(ai_response.content, output_base_name, config, logger)
            
            if saved_files:
                pbar.update(10)
                pbar.set_description("✅ 全部完成")
            else:
                logger.warning("未成功保存任何文件")
                return False

    except Exception as e:
        logger.error(f"处理过程中发生未预期的错误: {str(e)}", exc_info=True)
        return False

    total_duration = time.time() - start_time
    logger.info(f"="*50)
    logger.info(f"🎉 任务全部完成！总耗时: {total_duration:.2f}秒")
    logger.info(f"📁 归档文件保存在: {config['archive'].get('archive_root_dir', 'output_archive')}")
    logger.info(f"="*50)
    return True

# ================= 程序入口 =================
if __name__ == "__main__":
    # 1. 解析命令行参数
    args = parse_args()
    
    # 2. 加载配置
    config = load_config()
    
    # 3. 初始化日志
    logger = setup_logging(config)
    
    # 4. 打印欢迎信息
    print("\n" + "="*60)
    print("   硬件数据手册解析归档 Agent (hw_datasheet_agent)")
    print("="*60 + "\n")
    
    # 5. 处理输入文件路径
    input_dir = config['document'].get('input_dir', 'input_docs')
    # 命令行指定的文件优先级最高
    if args.file:
        input_file_name = args.file
    else:
        input_file_name = config['document']['input_file']
    
    # 拼接完整路径：如果是绝对路径直接用，否则拼接input_dir
    if os.path.isabs(input_file_name):
        input_file_path = input_file_name
    else:
        input_file_path = os.path.join(input_dir, input_file_name)
    
    # 6. 生成最终的输出基础名称（优先级：命令行 > 配置文件 > 自动生成）
    # 提取输入文件的纯名称（不带路径和后缀）
    input_file_basename = os.path.splitext(os.path.basename(input_file_path))[0]
    # 优先级1：命令行指定的输出名称
    if args.output:
        final_output_base_name = args.output.strip()
    # 优先级2：配置文件里指定的名称
    elif config['document'].get('output_base_name', '').strip():
        final_output_base_name = config['document']['output_base_name'].strip()
    # 优先级3：自动根据输入文件名生成
    else:
        final_output_base_name = f"{input_file_basename}_通信协议归档"
    
    logger.info(f"待解析文档: {input_file_path}")
    logger.info(f"输出文档基础名称: {final_output_base_name}")
    
    # 7. 运行主流程
    success = process_document(config, input_file_path, final_output_base_name, logger)
    sys.exit(0 if success else 1)